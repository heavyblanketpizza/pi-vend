"""Import seller report exports (CSV / XLSX) into the store tables.

Marketplace exports differ by report and change over time, so columns are
matched by Korean/English header synonyms instead of fixed layouts. The
header row is found automatically (exports often start with title rows),
CSV encodings UTF-8 and CP949 are both accepted, and re-importing a period
replaces that period's numbers instead of double counting.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from django.db import transaction

from .models import AdDaily, DailySales, DailyTraffic, DataImport, KeywordInflow, Product, Store

SYNONYMS: dict[str, list[str]] = {
    "date": ["날짜", "일자", "기간", "주문일", "결제일", "주문일자", "결제일자", "통계일", "date", "day"],
    "product": ["상품명", "노출상품명", "등록상품명", "옵션상품명", "상품", "product", "productname"],
    "orders": ["결제수", "주문수", "결제건수", "주문건수", "판매건수", "구매건수", "orders"],
    "units": ["결제상품수량", "판매수량", "주문수량", "결제수량", "수량", "units", "quantity"],
    "revenue": ["결제금액", "매출", "매출액", "판매금액", "결제액", "총매출", "순매출", "revenue", "sales", "gmv"],
    "visits": ["유입수", "방문수", "방문자수", "방문횟수", "세션수", "visits", "sessions", "visitors"],
    "page_views": ["페이지뷰", "조회수", "상품조회수", "pageviews", "pv"],
    "keyword": ["유입키워드", "검색어", "키워드", "검색키워드", "keyword", "searchterm", "query"],
    "impressions": ["노출수", "impressions"],
    "clicks": ["클릭수", "clicks"],
    "spend": ["집행광고비", "광고비", "총비용", "비용", "광고비용", "spend", "cost"],
    "conversions": ["총전환수", "전환수", "주문전환수", "conversions"],
    "ad_revenue": ["총전환매출액", "광고전환매출", "전환매출액", "전환매출", "conversionrevenue", "adrevenue"],
}

REQUIRED = {
    DataImport.Kind.SALES: [["date"], ["revenue", "orders"]],
    DataImport.Kind.TRAFFIC: [["date"], ["visits"]],
    DataImport.Kind.KEYWORDS: [["date"], ["keyword"], ["visits", "orders"]],
    DataImport.Kind.ADS: [["date"], ["spend"]],
}

FIELD_LABELS = {
    "date": "날짜", "product": "상품명", "orders": "주문수", "units": "판매수량", "revenue": "결제금액",
    "visits": "유입수", "page_views": "페이지뷰", "keyword": "키워드", "impressions": "노출수",
    "clicks": "클릭수", "spend": "광고비", "conversions": "전환수", "ad_revenue": "전환매출",
}


class ImportError_(ValueError):
    """Raised with a message the seller can act on."""


@dataclass
class ImportResult:
    kind: str
    rows: int
    first_date: date | None
    last_date: date | None
    columns: dict[str, str] = field(default_factory=dict)
    skipped: int = 0


def _norm(header: str) -> str:
    header = re.sub(r"\(.*?\)|\[.*?\]", "", str(header or ""))
    return re.sub(r"[\s_\-./:]+", "", header).lower()


def match_columns(headers: list[str]) -> dict[str, int]:
    """Map field -> column index. Exact synonym matches win over partial ones."""
    normalized = [_norm(h) for h in headers]
    found: dict[str, int] = {}
    used: set[int] = set()
    for exact in (True, False):
        for fld, names in SYNONYMS.items():
            if fld in found:
                continue
            for name in names:
                n = _norm(name)
                hit = next(
                    (i for i, h in enumerate(normalized) if i not in used and h and (h == n if exact else n in h)),
                    None,
                )
                if hit is not None:
                    found[fld] = hit
                    used.add(hit)
                    break
    return found


def detect_kind(columns: dict[str, int]) -> str | None:
    has = columns.__contains__
    if has("spend") and (has("clicks") or has("impressions")):
        return DataImport.Kind.ADS
    if has("keyword") and (has("visits") or has("orders")):
        return DataImport.Kind.KEYWORDS
    if has("revenue") or (has("orders") and not has("visits")):
        return DataImport.Kind.SALES
    if has("visits"):
        return DataImport.Kind.TRAFFIC
    return None


def _read_rows(filename: str, data: bytes) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise ImportError_(f"엑셀 파일을 읽을 수 없어요: {exc}") from exc
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("파일 인코딩을 알 수 없어요. UTF-8 또는 CP949(엑셀 기본) CSV로 저장해 주세요.")
    return list(csv.reader(io.StringIO(text)))


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]]:
    best = (-1, {})
    for i, row in enumerate(rows[:15]):
        cols = match_columns([str(c or "") for c in row])
        if len(cols) > len(best[1]):
            best = (i, cols)
    if best[0] < 0 or "date" not in best[1]:
        raise ImportError_("날짜 열을 찾지 못했어요. '날짜' 또는 '일자' 열이 있는 리포트를 올려 주세요.")
    return best


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value or "").strip()
    match = re.search(r"(\d{4})\D?(\d{1,2})\D?(\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match[1]), int(match[2]), int(match[3]))
    except ValueError:
        return None


def parse_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d.\-]", "", str(value))
    try:
        return float(text) if text not in {"", "-", "."} else 0.0
    except ValueError:
        return 0.0


def import_report(store: Store, filename: str, data: bytes, kind: str | None = None) -> ImportResult:
    rows = _read_rows(filename, data)
    header_index, columns = _find_header(rows)
    kind = kind or detect_kind(columns)
    if kind not in REQUIRED:
        raise ImportError_("어떤 리포트인지 알 수 없어요. 종류를 직접 선택해 주세요.")
    for options in REQUIRED[kind]:
        if not any(o in columns for o in options):
            names = " 또는 ".join(FIELD_LABELS[o] for o in options)
            raise ImportError_(f"{DataImport.Kind(kind).label} 리포트에 필요한 '{names}' 열이 없어요.")

    def cell(row, fld):
        i = columns.get(fld)
        return row[i] if i is not None and i < len(row) else None

    totals: dict[tuple, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    skipped = 0
    for row in rows[header_index + 1 :]:
        day = parse_date(cell(row, "date"))
        if day is None:
            skipped += 1  # totals rows, blank lines, footers
            continue
        if kind == DataImport.Kind.SALES:
            key = (day, str(cell(row, "product") or "").strip())
            fields = ("orders", "units", "revenue")
        elif kind == DataImport.Kind.TRAFFIC:
            key = (day,)
            fields = ("visits", "page_views")
        elif kind == DataImport.Kind.KEYWORDS:
            keyword = str(cell(row, "keyword") or "").strip()
            if not keyword:
                skipped += 1
                continue
            key = (day, keyword[:100])
            fields = ("visits", "orders", "revenue")
        else:
            key = (day, str(cell(row, "keyword") or "").strip()[:100])
            fields = ("impressions", "clicks", "spend", "conversions", "ad_revenue")
        for fld in fields:
            totals[key][fld] += parse_number(cell(row, fld))

    if not totals:
        raise ImportError_("가져올 데이터 행이 없어요.")

    with transaction.atomic():
        _store_rows(store, kind, totals)
    days = [k[0] for k in totals]
    result = ImportResult(
        kind=kind, rows=len(totals), first_date=min(days), last_date=max(days), skipped=skipped,
        columns={FIELD_LABELS[f]: str(rows[header_index][i]) for f, i in columns.items()},
    )
    DataImport.objects.create(
        store=store, kind=kind, filename=filename[:255], rows=result.rows,
        first_date=result.first_date, last_date=result.last_date,
    )
    return result


def _store_rows(store: Store, kind: str, totals: dict) -> None:
    n = lambda v: int(round(v))  # noqa: E731
    if kind == DataImport.Kind.SALES:
        names = {k[1] for k in totals if k[1]}
        products = {p.name: p for p in Product.objects.filter(store=store, name__in=names)}
        for name in names - products.keys():
            products[name] = Product.objects.create(store=store, name=name[:200])
        for (day, name), v in totals.items():
            DailySales.objects.update_or_create(
                store=store, product=products.get(name), date=day,
                defaults={"orders": n(v["orders"]), "units": n(v["units"]), "revenue": n(v["revenue"])},
            )
    elif kind == DataImport.Kind.TRAFFIC:
        for (day,), v in totals.items():
            DailyTraffic.objects.update_or_create(
                store=store, date=day, defaults={"visits": n(v["visits"]), "page_views": n(v["page_views"])}
            )
    elif kind == DataImport.Kind.KEYWORDS:
        for (day, keyword), v in totals.items():
            KeywordInflow.objects.update_or_create(
                store=store, date=day, keyword=keyword,
                defaults={"visits": n(v["visits"]), "orders": n(v["orders"]), "revenue": n(v["revenue"])},
            )
    else:
        for (day, keyword), v in totals.items():
            AdDaily.objects.update_or_create(
                store=store, date=day, keyword=keyword,
                defaults={
                    "impressions": n(v["impressions"]), "clicks": n(v["clicks"]), "spend": n(v["spend"]),
                    "conversions": n(v["conversions"]), "revenue": n(v["ad_revenue"]),
                },
            )
